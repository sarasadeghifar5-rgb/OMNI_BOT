"""
Document and file parser for extracting readable text from PDF, DOCX, TXT, Code, and data files.
"""
import io
import logging
from typing import Tuple
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

logger = logging.getLogger(__name__)

def extract_text_from_pdf(file_bytes: bytes, max_pages: int = 100) -> Tuple[str, int]:
    """
    Extracts text from PDF bytes.
    Returns (extracted_text, page_count).
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        total_pages = len(reader.pages)
        pages_to_read = min(total_pages, max_pages)
        
        extracted_chunks = []
        for i in range(pages_to_read):
            page_text = reader.pages[i].extract_text() or ""
            if page_text.strip():
                extracted_chunks.append(f"--- صفحه {i+1} ---\n{page_text.strip()}")
                
        full_text = "\n\n".join(extracted_chunks)
        return full_text, total_pages
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return "", 0

def extract_images_from_pdf(file_bytes: bytes, max_images: int = 5) -> list:
    """Extracts image bytes from PDF pages for scanned documents and math exam sheets."""
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        images = []
        for page in reader.pages:
            for img in page.images:
                images.append(img.data)
                if len(images) >= max_images:
                    return images
        return images
    except Exception as e:
        logger.error(f"Error extracting images from PDF: {e}")
        return []

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX bytes."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return ""

def extract_text_from_plain(file_bytes: bytes) -> str:
    """Extracts plain text with encoding fallbacks (UTF-8, Persian Windows-1256, Latin-1)."""
    for enc in ["utf-8", "utf-8-sig", "cp1256", "windows-1256", "latin-1"]:
        try:
            return file_bytes.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="ignore")

def parse_uploaded_file(file_name: str, file_bytes: bytes) -> Tuple[str, str, int]:
    """
    Parses uploaded file and returns (file_type, content_or_text, page_or_item_count).
    Types: 'pdf', 'docx', 'text', 'image', 'unsupported'
    """
    fn_lower = file_name.lower()

    if fn_lower.endswith(".pdf"):
        text, pages = extract_text_from_pdf(file_bytes)
        return "pdf", text, pages

    elif fn_lower.endswith((".docx", ".doc")):
        text = extract_text_from_docx(file_bytes)
        return "docx", text, 1

    elif fn_lower.endswith((".txt", ".py", ".js", ".ts", ".html", ".css", ".json", ".csv", ".xml", ".yaml", ".yml", ".sql", ".sh", ".md", ".log", ".ini", ".env")):
        text = extract_text_from_plain(file_bytes)
        return "text", text, 1

    elif fn_lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp")):
        return "image", "", 1

    else:
        # Try reading as plain text if small (< 1MB)
        if len(file_bytes) < 1024 * 1024:
            try:
                text = extract_text_from_plain(file_bytes)
                if any(c.isprintable() for c in text[:100]):
                    return "text", text, 1
            except Exception:
                pass
        return "unsupported", "", 0
