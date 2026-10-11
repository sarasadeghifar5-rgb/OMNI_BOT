import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

# Persian Zero-Width Non-Joiner (Half-Space)
NJ = "\u200c"

def create_accounting_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Accounting & Audit Professional Palette
    NAVY_DEEP = RGBColor(11, 25, 44)       # #0B192C Royal Deep Navy
    SLATE_DARK = RGBColor(30, 41, 59)      # #1E293B
    BG_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC Soft Off-White
    CARD_BG = RGBColor(255, 255, 255)      # Pure White Card
    CARD_BORDER = RGBColor(226, 232, 240)  # Clean Silver-Gray Border
    
    # Financial Accent Colors
    ACCENT_EMERALD = RGBColor(5, 150, 105) # #059669 Accounting Green (Balance/Assets)
    ACCENT_TEAL = RGBColor(13, 148, 136)   # #0D9488 Professional Teal
    ACCENT_BLUE = RGBColor(2, 132, 199)    # #0284C7 Trust Blue
    ACCENT_GOLD = RGBColor(217, 119, 6)    # #D97706 Audit Gold
    ACCENT_ROSE = RGBColor(225, 29, 72)    # #E11D48 Audit Alert Red
    ACCENT_INDIGO = RGBColor(79, 70, 229)  # #4F46E5 Analytical Indigo
    
    TEXT_DARK = RGBColor(15, 23, 42)       # Charcoal Dark
    TEXT_MUTED = RGBColor(71, 85, 105)     # Muted Slate
    TEXT_LIGHT = RGBColor(255, 255, 255)   # Crisp White
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184)
    
    # Specified Iranian Fonts
    FONT_TITR = "B Titr"
    FONT_NAZANIN = "B Nazanin"

    def apply_fa_text(paragraph, text="", font_name=FONT_NAZANIN, font_size=Pt(13), bold=False, color=TEXT_DARK, align=PP_ALIGN.RIGHT, space_after=Pt(4)):
        paragraph.alignment = align
        if space_after:
            paragraph.space_after = space_after
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set('rtl', '1')
        if align == PP_ALIGN.RIGHT:
            pPr.set('algn', 'r')
        elif align == PP_ALIGN.CENTER:
            pPr.set('algn', 'ctr')
        elif align == PP_ALIGN.LEFT:
            pPr.set('algn', 'l')
            
        if text:
            run = paragraph.add_run()
            run.text = text
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = bold
            run.font.color.rgb = color
            
            rPr = run._r.get_or_add_rPr()
            rPr.set('lang', 'fa-IR')
            rPr.set('altLang', 'en-US')
            cs = parse_xml(f'<a:cs xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font_name}"/>')
            rPr.append(cs)
            ea = parse_xml(f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font_name}"/>')
            rPr.append(ea)
            latin = parse_xml(f'<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font_name}"/>')
            rPr.append(latin)
            return run
        return None

    def add_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_accounting_header(slide, badge_text, title_text, subtitle_text=""):
        # Top ribbon container
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(1.3))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        
        # Category Badge
        p_badge = tf.paragraphs[0]
        apply_fa_text(p_badge, f"« {badge_text} »", font_name=FONT_TITR, font_size=Pt(11), bold=True, color=ACCENT_EMERALD, space_after=Pt(3))
        
        # Main Title (B Titr - Bold)
        p_title = tf.add_paragraph()
        apply_fa_text(p_title, title_text, font_name=FONT_TITR, font_size=Pt(21), bold=True, color=TEXT_DARK, space_after=Pt(3))
        
        # Subtitle (B Nazanin)
        if subtitle_text:
            p_sub = tf.add_paragraph()
            apply_fa_text(p_sub, subtitle_text, font_name=FONT_NAZANIN, font_size=Pt(12), bold=False, color=TEXT_MUTED, space_after=Pt(0))

    def add_accounting_footer(slide, slide_num, total_slides=17):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        tf = footer_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        pPr = p._p.get_or_add_pPr()
        pPr.set('rtl', '1')
        pPr.set('algn', 'r')
        
        r1 = p.add_run()
        r1.text = f"دانشجو: فاطمه عیوضی  |  استاد ارجمند: استاد بهرامیان  |  مفاهیم حسابرسی و رویدادهای مالی                       اسلاید {slide_num} از {total_slides}"
        r1.font.name = FONT_NAZANIN
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_MUTED
        rPr = r1._r.get_or_add_rPr()
        rPr.set('lang', 'fa-IR')
        rPr.set('altLang', 'en-US')

    def create_accounting_card(slide, left, top, width, height, accent_color=ACCENT_EMERALD, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.2)
        else:
            card.line.fill.background()
            
        strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.12))
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent_color
        strip.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: Title Slide (Elite Accounting Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1, NAVY_DEEP)
    
    # Frame Container
    frame = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.7), Inches(11.333), Inches(6.1))
    frame.fill.solid()
    frame.fill.fore_color.rgb = SLATE_DARK
    frame.line.color.rgb = RGBColor(51, 65, 85)
    frame.line.width = Pt(1.5)

    # Top Department Badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(1.05), Inches(4.7), Inches(0.48))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(6, 78, 59)
    badge.line.color.rgb = ACCENT_EMERALD
    badge.line.width = Pt(1.2)
    p_b = badge.text_frame.paragraphs[0]
    apply_fa_text(p_b, f"گروه حسابداری و حسابرسی | سمینار علمی کلاسی", font_name=FONT_TITR, font_size=Pt(12), bold=True, color=RGBColor(52, 211, 153), align=PP_ALIGN.CENTER)

    # Main Presentation Titles
    tb_title = s1.shapes.add_textbox(Inches(1.5), Inches(1.75), Inches(10.333), Inches(2.4))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True
    
    p_t1 = tf1.paragraphs[0]
    apply_fa_text(p_t1, f"مفهوم رویدادهای بعد از تاریخ ترازنامه", font_name=FONT_TITR, font_size=Pt(32), bold=True, color=TEXT_LIGHT, align=PP_ALIGN.CENTER, space_after=Pt(6))
    
    p_t2 = tf1.add_paragraph()
    apply_fa_text(p_t2, f"و وظایف حسابرس نسبت به آن", font_name=FONT_TITR, font_size=Pt(29), bold=True, color=RGBColor(56, 189, 248), align=PP_ALIGN.CENTER, space_after=Pt(8))
    
    p_t3 = tf1.add_paragraph()
    apply_fa_text(p_t3, f"بررسی استانداردهای حسابرسی، رویدادهای تعدیلی و غیرتعدیلی و رویه{NJ}های گزارشگری مالی", font_name=FONT_NAZANIN, font_size=Pt(15), bold=False, color=TEXT_LIGHT_MUTED, align=PP_ALIGN.CENTER)

    # Presenter and Professor Showcase Box
    info_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(4.45), Inches(9.333), Inches(1.85))
    info_box.fill.solid()
    info_box.fill.fore_color.rgb = NAVY_DEEP
    info_box.line.color.rgb = ACCENT_GOLD
    info_box.line.width = Pt(1.5)
    
    tf_info = info_box.text_frame
    tf_info.word_wrap = True
    
    p_prof = tf_info.paragraphs[0]
    apply_fa_text(p_prof, f"استاد گرانقدر: جناب استاد بهرامیان", font_name=FONT_TITR, font_size=Pt(17), bold=True, color=RGBColor(253, 224, 71), align=PP_ALIGN.CENTER, space_after=Pt(4))
    
    p_stu = tf_info.add_paragraph()
    apply_fa_text(p_stu, f"تهیه و ارائه: فاطمه عیوضی", font_name=FONT_TITR, font_size=Pt(17), bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER, space_after=Pt(4))
    
    p_course = tf_info.add_paragraph()
    apply_fa_text(p_course, f"درس: مباحث جاری در حسابداری و حسابرسی پیشرفته", font_name=FONT_NAZANIN, font_size=Pt(13.5), bold=False, color=TEXT_LIGHT_MUTED, align=PP_ALIGN.CENTER)

    # ==========================================
    # SLIDE 2: Table of Contents / Agenda
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2, NAVY_DEEP)
    
    h2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.2))
    tf2 = h2.text_frame
    p = tf2.paragraphs[0]
    apply_fa_text(p, f"« نقشه راه سمینار »", font_name=FONT_TITR, font_size=Pt(11), bold=True, color=ACCENT_EMERALD, space_after=Pt(2))
    p2 = tf2.add_paragraph()
    apply_fa_text(p2, f"فهرست مطالب و ساختار تخصصی ارائه", font_name=FONT_TITR, font_size=Pt(23), bold=True, color=TEXT_LIGHT)

    agenda_items = [
        (f"محور اول: مقدمه و اصل افشای کامل", f"تنوع نیازهای اطلاعاتی ذینفعان و محدودیت{NJ}های محیطی حسابداری", ACCENT_EMERALD),
        (f"محور دوم: تعریف و خط زمانی رویدادها", f"مفهوم رویدادهای پس از تاریخ ترازنامه و بازه زمانی وقوع", ACCENT_BLUE),
        (f"محور سوم: طبقه{NJ}بندی رویدادها (بیانیه ۲۵)", f"رویدادهای تعدیلی، رویدادهای غیرتعدیلی و رویدادهای محیطی", ACCENT_TEAL),
        (f"محور چهارم: رویدادهای احتمالی و صورت{NJ}های فرضی", f"بررسی مقتضیات احتمالی و ارائه صورت{NJ}های مالی فرضی [Pro-Forma]", ACCENT_GOLD),
        (f"محور پنجم: روش{NJ}های رسیدگی در ۳ مقطع زمانی", f"رسیدگی تا تاریخ گزارش، پیش از انتشار و پس از انتشار عمومی", ACCENT_ROSE),
        (f"محور ششم: تاریخ{NJ}گذاری دوگانه و جمع{NJ}بندی", f"مفهوم Dual Dating، تحدید مسئولیت حسابرس و ماتریس نهایی", ACCENT_INDIGO)
    ]
    
    card_w = Inches(3.7)
    card_h = Inches(2.2)
    positions = [
        (Inches(8.8), Inches(1.8)),  # Col 1 (Rightmost)
        (Inches(4.8), Inches(1.8)),  # Col 2
        (Inches(0.8), Inches(1.8)),  # Col 3
        (Inches(8.8), Inches(4.3)),  # Row 2 Col 1
        (Inches(4.8), Inches(4.3)),  # Row 2 Col 2
        (Inches(0.8), Inches(4.3))   # Row 2 Col 3
    ]
    
    for idx, (title, desc, acc) in enumerate(agenda_items):
        pos_x, pos_y = positions[idx]
        cd = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x, pos_y, card_w, card_h)
        cd.fill.solid()
        cd.fill.fore_color.rgb = SLATE_DARK
        cd.line.color.rgb = RGBColor(51, 65, 85)
        cd.line.width = Pt(1)
        
        st = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x, pos_y, card_w, Inches(0.12))
        st.fill.solid()
        st.fill.fore_color.rgb = acc
        st.line.fill.background()
        
        tf_c = cd.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = Inches(0.2)
        p_c1 = tf_c.paragraphs[0]
        apply_fa_text(p_c1, title, font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_LIGHT, space_after=Pt(6))
        
        p_c2 = tf_c.add_paragraph()
        apply_fa_text(p_c2, desc, font_name=FONT_NAZANIN, font_size=Pt(12), bold=False, color=TEXT_LIGHT_MUTED)

    add_accounting_footer(s2, 2)

    # ==========================================
    # SLIDE 3: Introduction & Full Disclosure
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3, BG_LIGHT)
    add_accounting_header(s3, "بخش اول: مبانی نظری گزارشگری مالی", "مقدمه؛ استفاده‌کنندگان گزارش‌های مالی و اصل افشای کامل", "ضرورت ارزیابی شفاف عملکرد و افشای رویدادهای بااهمیت در حسابداری")
    
    w3 = Inches(3.7)
    h3 = Inches(4.8)
    
    # Card 1 (Rightmost)
    c1 = create_accounting_card(s3, Inches(8.8), Inches(1.8), w3, h3, ACCENT_BLUE)
    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۱. تنوع نیازهای اطلاعاتی ذینفعان", font_name=FONT_TITR, font_size=Pt(14), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- استفاده{NJ}کنندگان صورت{NJ}های مالی از اقشار مختلف اجتماع با هدف{NJ}ها و سطوح آگاهی گوناگون تشکیل شده{NJ}اند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- نیازهای اطلاعاتی و علایق سرمایه{NJ}گذاران، بستانکاران و مدیران بسیار متفاوت است.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- ارضای تمام نیازهای فردی امکان{NJ}پذیر نبوده و گزارشگری عمومی استاندارد ضرورت دارد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK)

    # Card 2 (Center)
    c2 = create_accounting_card(s3, Inches(4.8), Inches(1.8), w3, h3, ACCENT_GOLD)
    tf = c2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۲. محدودیت‌های محیطی و عدم‌قطعیت", font_name=FONT_TITR, font_size=Pt(14), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- حسابداران در تهیه اطلاعات مالی با محدودیت{NJ}های محیطی و فرض{NJ}های بنیادی حسابداری مواجهند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- تقریباً کلیه ارقام صورت{NJ}های مالی به دلیل ابهامات آتی، دارای ماهیت غیرقطعی هستند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- اتکای صورت{NJ}های مالی بر برآوردها، نیاز به داده{NJ}های تکمیلی بعدی را ضروری می{NJ}سازد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK)

    # Card 3 (Leftmost)
    c3 = create_accounting_card(s3, Inches(0.8), Inches(1.8), w3, h3, ACCENT_EMERALD)
    tf = c3.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۳. اصل بنیادی افشای کامل", font_name=FONT_TITR, font_size=Pt(14), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- یکی از اصول اولیه حسابداری، افشای تمامی واقعیت{NJ}های بااهمیت در زمینه رویدادها و فعالیت{NJ}های مالی است.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"اهداف کلیدی اصل افشای کامل:", font_name=FONT_TITR, font_size=Pt(12), bold=True, color=ACCENT_EMERALD, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"• امکان ارزیابی عملکرد واقعی واحد تجاری\n• داوری در مورد چگونگی استفاده از منابع موجود\n• پیش{NJ}بینی دورنمای آینده و ریسک{NJ}های پیش{NJ}رو", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    add_accounting_footer(s3, 3)

    # ==========================================
    # SLIDE 4: Definition & Timeline
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4, BG_LIGHT)
    add_accounting_header(s4, "بخش اول: مبانی نظری گزارشگری مالی", "تعریف رویدادهای پس از تاریخ ترازنامه و خط زمانی وقوع", "بررسی اثر گذشت زمان بر ابهامات و ارقام منعکس در صورت‌های مالی")
    
    # Top Card: Definition
    c_def = create_accounting_card(s4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(2.35), ACCENT_BLUE)
    tf_def = c_def.text_frame
    tf_def.word_wrap = True
    tf_def.margin_left = tf_def.margin_right = Inches(0.3)
    tf_def.margin_top = Inches(0.25)
    
    p = tf_def.paragraphs[0]
    apply_fa_text(p, f"تعریف استانداردی رویدادهای پس از تاریخ ترازنامه [Subsequent Events]:", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf_def.add_paragraph()
    apply_fa_text(p, f"«به آن گروه از رویدادها که پس از تاریخ ترازنامه، اما قبل از صدور و انتشار صورت{NJ}های مالی و گزارش حسابرسی اتفاق می{NJ}افتد و دارای آثار بااهمیت بر صورت{NJ}های مالی بوده و نیاز به اصلاح یا افشا در صورت{NJ}های مالی دارد، رویدادهای پس از تاریخ ترازنامه گفته می{NJ}شود.»", font_name=FONT_NAZANIN, font_size=Pt(13.5), bold=False, color=RGBColor(30, 41, 59), space_after=Pt(6))
    p = tf_def.add_paragraph()
    apply_fa_text(p, f"نکته کلیدی: با گذشت زمان و دستیابی به اطلاعات اضافی، بسیاری از ابهامات برطرف می{NJ}شود؛ لذا این رویدادها بر اعتبار صورت{NJ}های مالی و تصمیمات استفاده{NJ}کنندگان شدیداً اثرگذار است.", font_name=FONT_NAZANIN, font_size=Pt(12.5), bold=True, color=ACCENT_BLUE)

    # Bottom Area: 3 Timeline Stages
    timeline_boxes = [
        (f"مرحله اول: پایان سال مالی", f"تاریخ ترازنامه (۲۹ اسفند)\nاندازه{NJ}گیری منابع و روابط مالی در پایان دوره مورد گزارش", ACCENT_INDIGO, Inches(8.8)),
        (f"مرحله دوم: دوره رسیدگی میدانی", f"دوره پس از تاریخ ترازنامه\nفاصله زمانی تا تاریخ گزارش حسابرس جهت تکمیل آزمون{NJ}ها", ACCENT_GOLD, Inches(4.8)),
        (f"مرحله سوم: انتشار و ارائه عمومی", f"تاریخ انتشار صورت{NJ}های مالی\nتصویب و ارائه عمومی گزارش حسابرسی به مجامع و سهامداران", ACCENT_EMERALD, Inches(0.8))
    ]
    
    for title, desc, col, pos_x in timeline_boxes:
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x, Inches(4.45), Inches(3.7), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = col
        box.line.width = Pt(1.5)
        
        st = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x, Inches(4.45), Inches(3.7), Inches(0.12))
        st.fill.solid()
        st.fill.fore_color.rgb = col
        st.line.fill.background()
        
        tf_b = box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_right = Inches(0.2)
        tf_b.margin_top = Inches(0.2)
        
        p = tf_b.paragraphs[0]
        apply_fa_text(p, title, font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
        p = tf_b.add_paragraph()
        apply_fa_text(p, desc, font_name=FONT_NAZANIN, font_size=Pt(12), bold=False, color=TEXT_MUTED)

    add_accounting_footer(s4, 4)

    # ==========================================
    # SLIDE 5: Classification by Statement No. 25
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5, BG_LIGHT)
    add_accounting_header(s5, "بخش دوم: طبقه‌بندی رویدادها", "طبقه‌بندی رویدادها بر اساس بیانیه روش‌های حسابرسی شماره ۲۵", "دسته‌بندی تخصصی معاملات و رویدادهای پس از تاریخ ترازنامه به سه گروه مستقل")
    
    c1 = create_accounting_card(s5, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8), ACCENT_EMERALD)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = Inches(0.25)
    tf1.margin_top = Inches(0.25)
    p = tf1.paragraphs[0]
    apply_fa_text(p, f"رویدادهای گروه اول\n[دارای اثر مستقیم بر ارقام]", font_name=FONT_TITR, font_size=Pt(14), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf1.add_paragraph()
    apply_fa_text(p, f"- رویدادهایی است که مستقیماً بر صورت{NJ}های مالی تأثیر می{NJ}گذارند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf1.add_paragraph()
    apply_fa_text(p, f"- اطلاعات بعدی مبنایی دقیق{NJ}تر برای برآوردها یا پیش{NJ}بینی{NJ}ها فراهم می{NJ}آورد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf1.add_paragraph()
    apply_fa_text(p, f"- الزام قطعی: اگر این اطلاعات در تاریخ ترازنامه موجود بود استفاده می{NJ}شد؛ لذا باید اصلاحات مناسب در صورت{NJ}های مالی انجام شود.", font_name=FONT_NAZANIN, font_size=Pt(12.5), bold=True, color=ACCENT_EMERALD)

    c2 = create_accounting_card(s5, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8), ACCENT_BLUE)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = Inches(0.25)
    tf2.margin_top = Inches(0.25)
    p = tf2.paragraphs[0]
    apply_fa_text(p, f"رویدادهای گروه دوم\n[فاقد اثر مستقیم بر ارقام]", font_name=FONT_TITR, font_size=Pt(14), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf2.add_paragraph()
    apply_fa_text(p, f"- رویدادهایی است که تأثیر مستقیم بر صورت{NJ}های مالی ندارند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf2.add_paragraph()
    apply_fa_text(p, f"- اصلاح صورت{NJ}های مالی دوره{NJ}های قبل را الزامی نمی{NJ}کنند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf2.add_paragraph()
    apply_fa_text(p, f"- ضرورت افشا: تأثیرشان ممکن است به گونه{NJ}ای باشد که افشای آن{NJ}ها در یادداشت{NJ}ها برای آگاهی ذینفعان کاملاً الزامی باشد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), bold=True, color=ACCENT_BLUE)

    c3 = create_accounting_card(s5, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8), ACCENT_ROSE)
    tf3 = c3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_right = Inches(0.25)
    tf3.margin_top = Inches(0.25)
    p = tf3.paragraphs[0]
    apply_fa_text(p, f"رویدادهای گروه سوم\n[رویدادهای غیرحسابداری]", font_name=FONT_TITR, font_size=Pt(14), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf3.add_paragraph()
    apply_fa_text(p, f"- رویدادهای غیرحسابداری نظیر جنگ، تغییرات مدیریتی، ورشکستگی مشتریان عمده و اعتصابات.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf3.add_paragraph()
    apply_fa_text(p, f"- افشای این اطلاعات معمولاً موجب شک و تردید نسبت به علت رخداد آن{NJ}ها می{NJ}شود.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf3.add_paragraph()
    apply_fa_text(p, f"- چالش افشا: همان{NJ}طور که افشای آن{NJ}ها آگاهی{NJ}دهنده است، می{NJ}تواند موجب گمراهی استفاده{NJ}کنندگان گردد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), bold=True, color=ACCENT_ROSE)

    add_accounting_footer(s5, 5)

    # ==========================================
    # SLIDE 6: Group 1 - Adjusting Events
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6, BG_LIGHT)
    add_accounting_header(s6, "بخش دوم: طبقه‌بندی رویدادها", "گروه اول: رویدادهای تعدیلی [Adjusting Events]", "رویدادهای مستلزم اصلاح و تعدیل ارقام در متن صورت‌های مالی اساسی")
    
    w6 = Inches(5.7)
    h6 = Inches(4.8)
    
    c_adj_def = create_accounting_card(s6, Inches(6.8), Inches(1.8), w6, h6, ACCENT_EMERALD)
    tf = c_adj_def.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"مفهوم و ماهیت رویدادهای تعدیلی", font_name=FONT_TITR, font_size=Pt(15), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- رویدادهایی هستند که منجر به روشن شدن اطلاعات بیشتر درباره شرایطی می{NJ}شود که در تاریخ ترازنامه وجود داشته{NJ}اند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- علت رخداد: ناشی از نبود اطلاعات کافی طی دوره مالی است و با دسترسی به اطلاعات جدید، ارزیابی{NJ}ها و برآوردها تغییر می{NJ}کند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- اثر حسابداری: مستقیماً بر مبالغ منعکس در صورت{NJ}های مالی موثرند و صورت{NJ}های مالی قبل از انتشار باید به نحو مناسبی اصلاح شوند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), bold=True, color=ACCENT_EMERALD, space_after=Pt(8))
    p = tf.add_paragraph()
    apply_fa_text(p, f"نمونه‌های بارز رویدادهای تعدیلی:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"۱. ورشکستگی مشتری پس از پایان دوره مالی به دلیل بدهی{NJ}های قبلی\n۲. حل{NJ}وفصل قطعی دعاوی حقوقی معوق با رقم مشخص دادگاه", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_MUTED)

    c_adj_cond = create_accounting_card(s6, Inches(0.8), Inches(1.8), w6, h6, ACCENT_BLUE)
    tf = c_adj_cond.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"شرایط سه‌گانه انعکاس در متن صورت‌های مالی", font_name=FONT_TITR, font_size=Pt(15), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"آثار رویدادهای پس از تاریخ ترازنامه را هنگامی می{NJ}توان در متن صورت{NJ}های مالی منعکس کرد که واجد هر ۳ شرط زیر باشند:", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(8))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"شرط اول: وجود شواهد درباره شرایط تاریخی", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_BLUE, space_after=Pt(2))
    p = tf.add_paragraph()
    apply_fa_text(p, f"فراهم{NJ}کننده شواهد و اطلاعات اضافی درباره وضعیت{NJ}هایی باشد که در تاریخ ترازنامه وجود داشته است.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(6))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"شرط دوم: اثرگذاری بر برآوردهای گزارشگری", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_BLUE, space_after=Pt(2))
    p = tf.add_paragraph()
    apply_fa_text(p, f"روی برآوردهای اجتناب{NJ}ناپذیر در فرایند تهیه صورت{NJ}های مالی اثر مستقیم داشته باشد.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(6))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"شرط سوم: ضرورت اصلاح ارقام مالی اساسی", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_BLUE, space_after=Pt(2))
    p = tf.add_paragraph()
    apply_fa_text(p, f"تعدیل اطلاعات مالی صورت{NJ}های مالی اساسی ناشی از برآوردها را کاملاً ضروری سازد.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    add_accounting_footer(s6, 6)

    # ==========================================
    # SLIDE 7: Group 2 - Non-Adjusting Events
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7, BG_LIGHT)
    add_accounting_header(s7, "بخش دوم: طبقه‌بندی رویدادها", "گروه دوم: رویدادهای غیرتعدیلی [Non-Adjusting Events]", "رویدادهای ناشی از شرایط پس از تاریخ ترازنامه و ضوابط افشا در یادداشت‌ها")
    
    c_non_def = create_accounting_card(s7, Inches(6.8), Inches(1.8), w6, h6, ACCENT_GOLD)
    tf = c_non_def.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"مفهوم و ماهیت رویدادهای غیرتعدیلی", font_name=FONT_TITR, font_size=Pt(15), bold=True, color=TEXT_DARK, space_after=Pt(8))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- رویدادهایی است که منجر به روشن شدن اطلاعات بیشتر درباره شرایطی می{NJ}شود که در تاریخ ترازنامه وجود نداشته و پس از آن اتفاق افتاده است.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- عدم تعدیل ارقام: رویدادهایی هستند که مستلزم اصلاح و تعدیل ارقام مالی نبوده و ارقام ترازنامه را تغییر نمی{NJ}دهند.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- ضرورت افشا: صرفاً افشای آن{NJ}ها از راه یادداشت{NJ}های پیوست کافی خواهد بود تا استفاده{NJ}کننده تصویر درست و منصفانه{NJ}ای به دست آورد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), bold=True, color=ACCENT_GOLD, space_after=Pt(8))
    p = tf.add_paragraph()
    apply_fa_text(p, f"نمونه‌های متداول رویدادهای غیرتعدیلی:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"• خسارات آتش{NJ}سوزی در انبار یا کارخانه در سال مالی جدید\n• صدور سهام جدید یا انتشار اوراق مشارکت با مبالغ بااهمیت\n• کاهش ناگهانی و شدید ارزش بازار سرمایه{NJ}گذاری{NJ}ها", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_MUTED)

    c_non_cond = create_accounting_card(s7, Inches(0.8), Inches(1.8), w6, h6, ACCENT_INDIGO)
    tf = c_non_cond.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"ضوابط سه‌گانه افشا در یادداشت‌های همراه", font_name=FONT_TITR, font_size=Pt(15), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"رویدادهای پس از تاریخ ترازنامه را هنگامی می{NJ}توان از راه یادداشت{NJ}های همراه صورت{NJ}های مالی افشا کرد که:", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(8))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"۱. عدم ریشه در شرایط تاریخ ترازنامه", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_INDIGO, space_after=Pt(2))
    p = tf.add_paragraph()
    apply_fa_text(p, f"ناشی از شرایط موجود در تاریخ ترازنامه نباشد (واقعه{NJ}ای کاملاً جدید باشد).", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(6))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"۲. وقوع حتمی پس از پایان دوره مالی", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_INDIGO, space_after=Pt(2))
    p = tf.add_paragraph()
    apply_fa_text(p, f"حتماً و بدون تردید پس از تاریخ ترازنامه واقع شده باشد.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(6))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"۳. فقدان معیارهای اصلاح متن صورت‌ها", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_INDIGO, space_after=Pt(2))
    p = tf.add_paragraph()
    apply_fa_text(p, f"بدون شرایط لازم برای اصلاح ارقام اساسی باشد، اما افشای آن جهت جلوگیری از گمراهی ذینفعان ضروری باشد.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    add_accounting_footer(s7, 7)

    # ==========================================
    # SLIDE 8: Contingent Events & Pro-Forma Statements
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_background(s8, BG_LIGHT)
    add_accounting_header(s8, "بخش دوم: طبقه‌بندی رویدادها", "رویدادهای احتمالی و صورت‌های مالی فرضی [Pro-Forma]", "بررسی مقتضیات احتمالی و شیوه انعکاس رویدادهای دگرگون‌کننده ساختار سرمایه")
    
    c_cont = create_accounting_card(s8, Inches(6.8), Inches(1.8), w6, h6, ACCENT_BLUE)
    tf = c_cont.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"رویدادهای احتمالی [Contingent Events]", font_name=FONT_TITR, font_size=Pt(15), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"رویدادهای احتمالی به عنوان بخشی از رویدادهای پس از تاریخ ترازنامه است و باید چگونگی برخورد با آن{NJ}ها مورد توجه قرار گیرد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"تعریف استانداردی رویداد احتمالی:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_BLUE, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"«عبارت از وضعیت، شرایط یا مجموعه{NJ}ای از مقتضیات موجود است که با ضرایب احتمال مختلفی ممکن است بر اثر رویدادهای آتی مربوط به تحصیل یا از دست دادن دارایی یا ایجاد یا پرهیز از بدهی (همراه با کسب سود یا تحمل زیان) منجر شود.»", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"الزام حسابرس: سنجش ضرایب احتمال وقوع جهت تصمیم{NJ}گیری درباره شناسایی ذخیره یا افشا در یادداشت{NJ}ها.", font_name=FONT_NAZANIN, font_size=Pt(12), bold=True, color=TEXT_MUTED)

    c_prof = create_accounting_card(s8, Inches(0.8), Inches(1.8), w6, h6, ACCENT_GOLD)
    tf = c_prof.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"صورت‌های مالی فرضی [Pro-Forma]", font_name=FONT_TITR, font_size=Pt(15), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"گاه رویدادهای پس از تاریخ ترازنامه چنان بااهمیت است که باید همراه صورت{NJ}های مالی اصلی، صورت{NJ}های مالی فرضی ارائه شود.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"مفهوم و هدف ارائه فرضی:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_GOLD, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"تاثیر رویدادهای پس از تاریخ ترازنامه را به گونه{NJ}ای نشان می{NJ}دهد که گویی در تاریخ ترازنامه اتفاق افتاده است.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"قالب و موارد کاربرد در حسابداری:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=TEXT_DARK, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- صورت{NJ}های فرضی (معمولاً تنها ترازنامه) در یک ستون و در کنار صورت{NJ}های مالی اصلی ارائه می{NJ}شود.\n- این شکل از افشا تنها در مواردی مانند ادغام شرکت{NJ}ها [Mergers] به کار می{NJ}رود که ساختار دارایی یا سرمایه دگرگون شده است.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_MUTED)

    add_accounting_footer(s8, 8)

    # ==========================================
    # SLIDE 9: Audit Objectives
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_background(s9, BG_LIGHT)
    add_accounting_header(s9, "بخش سوم: مسئولیت‌ها و اهداف حسابرسی", "اهداف حسابرسی رویدادهای پس از تاریخ ترازنامه", "کسب اطمینان معقول از شناسایی، ثبت مناسب و صیانت از اعتبار گزارشگری مالی")
    
    grid_w = Inches(5.7)
    grid_h = Inches(2.25)
    
    b1 = create_accounting_card(s9, Inches(6.8), Inches(1.8), grid_w, grid_h, ACCENT_TEAL)
    tf = b1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۱. کسب اطمینان معقول [Reasonable Assurance]", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"کسب اطمینان معقول درباره شناسایی تمامی رویدادهای مهم پس از تاریخ ترازنامه و در صورت لزوم افشا یا به حساب گرفتن آن{NJ}ها در صورت{NJ}های مالی.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    b2 = create_accounting_card(s9, Inches(0.8), Inches(1.8), grid_w, grid_h, ACCENT_BLUE)
    tf = b2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۲. انعکاس و افشای مناسب [Appropriate Accounting]", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"تمام رویدادهای پس از تاریخ ترازنامه در صورت{NJ}های مالی شناسایی و به نحو مناسبی در حساب{NJ}ها منظور یا در یادداشت{NJ}ها افشا شده{NJ}اند.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    b3 = create_accounting_card(s9, Inches(6.8), Inches(4.35), grid_w, grid_h, ACCENT_EMERALD)
    tf = b3.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۳. صحت ارزش‌های ترازنامه‌ای [Valuation]", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"ارزش{NJ}های ترازنامه{NJ}ای به طور صحیح شامل آثار رویدادهای پس از تاریخ ترازنامه است و شواهد تکمیلی در ارزش دارایی{NJ}ها و بدهی{NJ}ها لحاظ شده است.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    b4 = create_accounting_card(s9, Inches(0.8), Inches(4.35), grid_w, grid_h, ACCENT_ROSE)
    tf = b4.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۴. صیانت از تمامیت صورت‌های مالی [Completeness]", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"حسابرس باید اطمینان حاصل کند که هیچ رویداد بااهمیتی که موجب مخدوش شدن نظر ارائه شده درباره تمامیت صورت{NJ}های مالی شود، وجود ندارد.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    add_accounting_footer(s9, 9)

    # ==========================================
    # SLIDE 10: Auditor Responsibilities Roadmap (3 Intervals)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_background(s10, BG_LIGHT)
    add_accounting_header(s10, "بخش چهارم: روش‌های رسیدگی حسابرسی", "مقاطع زمانی سه‌گانه وظایف حسابرس (نقشه راه رسیدگی)", "تفکیک زمانی مسئولیت‌های حسابرس در رهنمودهای حرفه‌ای حسابرسی")
    
    w10 = Inches(3.7)
    h10 = Inches(4.8)
    
    p1 = create_accounting_card(s10, Inches(8.8), Inches(1.8), w10, h10, ACCENT_EMERALD)
    tf = p1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"مقطع اول:\nتا تاریخ گزارش حسابرس", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"بازه: از تاریخ ترازنامه تا پایان عملیات میدانی", font_name=FONT_TITR, font_size=Pt(11), bold=True, color=ACCENT_EMERALD, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- نوع مسئولیت: مسئولیت فعالانه [Active]", font_name=FONT_TITR, font_size=Pt(12), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- حسابرس برای شناسایی رویدادها باید روش{NJ}هایی برای کسب شواهد کافی و قابل قبول طراحی و اجرا کند.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- شامل بررسی رویه{NJ}ها، صورتجلسات، صورت{NJ}های میان{NJ}دوره{NJ}ای، دفاتر و پرس{NJ}وجو از مدیریت.", font_name=FONT_NAZANIN, font_size=Pt(11.5), color=TEXT_MUTED)

    p2 = create_accounting_card(s10, Inches(4.8), Inches(1.8), w10, h10, ACCENT_GOLD)
    tf = p2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"مقطع دوم:\nپس از گزارش تا انتشار صورت‌ها", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"بازه: از امضای گزارش تا انتشار رسمی", font_name=FONT_TITR, font_size=Pt(11), bold=True, color=ACCENT_GOLD, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- نوع مسئولیت: مسئولیت واکنشی [Passive]", font_name=FONT_TITR, font_size=Pt(12), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- حسابرس مسئولیت تجسس فعال ندارد؛ مسئولیت آگاه کردن حسابرس به عهده مدیر واحد مورد رسیدگی است.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- در صورت آگاهی از حقایق، حسابرس مکلف به اقدام فوری (اصلاح یا اظهارنظر مشروط) است.", font_name=FONT_NAZANIN, font_size=Pt(11.5), color=TEXT_MUTED)

    p3 = create_accounting_card(s10, Inches(0.8), Inches(1.8), w10, h10, ACCENT_ROSE)
    tf = p3.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"مقطع سوم:\nپس از انتشار صورت‌های مالی", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"بازه: پس از انتشار عمومی و دردسترس بودن", font_name=FONT_TITR, font_size=Pt(11), bold=True, color=ACCENT_ROSE, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- نوع مسئولیت: حقایق کشف شده بعدی", font_name=FONT_TITR, font_size=Pt(12), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- حسابرس نسبت به انجام پرس{NJ}وجو مسئولیتی ندارد مگر حقایقی کشف شود که در تاریخ گزارش وجود داشته است.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- صدور گزارش تجدیدنظر شده با بند تأکید بر مطلب یا سلب اتکا از گزارش در صورت امتناع مدیریت.", font_name=FONT_NAZANIN, font_size=Pt(11.5), color=TEXT_MUTED)

    add_accounting_footer(s10, 10)

    # ==========================================
    # SLIDE 11: Audit Procedures up to Report Date
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_background(s11, BG_LIGHT)
    add_accounting_header(s11, "بخش چهارم: روش‌های رسیدگی حسابرسی", "مرحله اول: روش‌های رسیدگی تا تاریخ گزارش حسابرس", "اقدامات ضروری حسابرس در نزدیک‌ترین زمان ممکن به تاریخ گزارش حسابرس")
    
    b1 = create_accounting_card(s11, Inches(6.8), Inches(1.8), grid_w, grid_h, ACCENT_BLUE)
    tf = b1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۱. بررسی رویه‌های مدیریت", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"بررسی روش{NJ}های مدیریت برای حصول اطمینان از شناسایی رویدادهای پس از تاریخ ترازنامه و پایش رویه{NJ}های کنترل داخلی واحد مورد رسیدگی.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    b2 = create_accounting_card(s11, Inches(0.8), Inches(1.8), grid_w, grid_h, ACCENT_EMERALD)
    tf = b2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۲. مطالعه و بررسی صورتجلسات", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"مطالعه و بررسی صورتجلسات مجامع عمومی صاحبان سهام، هیئت مدیره و کمیته{NJ}های هیئت مدیره (مانند کمیته حسابرسی) و کاوش درباره جلسات نهایی.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    b3 = create_accounting_card(s11, Inches(6.8), Inches(4.35), grid_w, grid_h, ACCENT_GOLD)
    tf = b3.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۳. بررسی صورت‌های میان‌دوره‌ای و بودجه‌ها", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"مطالعه آخرین صورت{NJ}های مالی میان{NJ}دوره{NJ}ای و در صورت لزوم، بودجه، گردش وجوه نقد برآوردی و گزارش{NJ}های مرتبط مدیریت و مقایسه با ارقام حسابرسی{NJ}شده.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    b4 = create_accounting_card(s11, Inches(0.8), Inches(4.35), grid_w, grid_h, ACCENT_INDIGO)
    tf = b4.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    apply_fa_text(p, f"۴. بررسی ثبت‌های حسابداری و آزمون انقطاع", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"بررسی ثبت{NJ}های دفتر کل، روزنامه، دریافت{NJ}ها و پرداخت{NJ}های نقدی غیرعادی تا آخرین روز عملیات و اجرای آزمون انقطاع زمانی [Cut-off].", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    add_accounting_footer(s11, 11)

    # ==========================================
    # SLIDE 12: Management Inquiries and Representation Letter
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_background(s12, BG_LIGHT)
    add_accounting_header(s12, "بخش چهارم: روش‌های رسیدگی حسابرسی", "پرس‌وجوهای کلیدی از مدیریت و اخذ تأییدیه مدیران", "چک‌لیست پرس‌وجو از مدیریت و اهداف اخذ تأییدیه کتبی [Management Representation]")
    
    c_inq = create_accounting_card(s12, Inches(5.2), Inches(1.8), Inches(7.333), Inches(4.8), ACCENT_TEAL)
    tf_inq = c_inq.text_frame
    tf_inq.word_wrap = True
    tf_inq.margin_left = tf_inq.margin_right = Inches(0.3)
    tf_inq.margin_top = Inches(0.2)
    
    p = tf_inq.paragraphs[0]
    apply_fa_text(p, f"محورهای اصلی پرس‌وجو از مدیریت واحد مورد رسیدگی:", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    
    inquiries_text = [
        f"- وجود هرگونه تعهدات یا بدهی{NJ}های احتمالی بااهمیت در تاریخ حسابرسی",
        f"- ایجاد تغییرات قابل توجه در سرمایه، بدهی{NJ}های بلندمدت یا سرمایه در گردش",
        f"- اصلاحات غیرعادی یا تغییرات در برآوردهای حسابداری پس از تاریخ ترازنامه",
        f"- تغییر در روند فروش{NJ}ها، هزینه{NJ}ها، سود و قیمت{NJ}های مواد خام و ارزش جایگزینی",
        f"- لغو سفارشات فروش یا متضرر شدن مشتریان اصلی پس از تاریخ حسابرسی",
        f"- رویدادهای پیش{NJ}بینی{NJ}نشده (سرقت{NJ}ها، آتش{NJ}سوزی{NJ}ها، انفجارها و غیره)",
        f"- فروش دارایی{NJ}های غیرجاری، صدور سهام یا اوراق مشارکت و ادغام یا انحلال",
        f"- معاملات با اشخاص وابسته و رویدادهای مخدوش{NJ}کننده «فرض تداوم فعالیت»"
    ]
    for inq in inquiries_text:
        p = tf_inq.add_paragraph()
        apply_fa_text(p, inq, font_name=FONT_NAZANIN, font_size=Pt(11.5), color=TEXT_DARK, space_after=Pt(3))

    c_rep = create_accounting_card(s12, Inches(0.8), Inches(1.8), Inches(4.1), Inches(4.8), ACCENT_GOLD)
    tf_rep = c_rep.text_frame
    tf_rep.word_wrap = True
    tf_rep.margin_left = tf_rep.margin_right = Inches(0.25)
    tf_rep.margin_top = Inches(0.25)
    
    p = tf_rep.paragraphs[0]
    apply_fa_text(p, f"تأییدیه‌های مدیریت\n[Management Representation]", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf_rep.add_paragraph()
    apply_fa_text(p, f"- اخذ تأییدیه به امضای مدیر با دو هدف راهبردی صورت می{NJ}گیرد:", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(4))
    p = tf_rep.add_paragraph()
    apply_fa_text(p, f"۱. تأکید بر اینکه مسئولیت صورت{NJ}های مالی به عهده مدیریت است.\n۲. تهیه شواهد و مدارک حسابرسی برای مواردی که روش{NJ}های معمولی حسابرسی کفایت نمی{NJ}کند.", font_name=FONT_NAZANIN, font_size=Pt(12), bold=True, color=ACCENT_GOLD, space_after=Pt(8))
    
    p = tf_rep.add_paragraph()
    apply_fa_text(p, f"حسابرسان شعب و اجزای فرعی:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=TEXT_DARK, space_after=Pt(3))
    p = tf_rep.add_paragraph()
    apply_fa_text(p, f"حسابرس اصلی باید روش{NJ}های حسابرسان شعب را درباره رویدادهای پس از تاریخ ترازنامه بررسی و لزوم آگاه کردن آنان را از تاریخ گزارش مدنظر قرار دهد.", font_name=FONT_NAZANIN, font_size=Pt(11.5), color=TEXT_MUTED)

    add_accounting_footer(s12, 12)

    # ==========================================
    # SLIDE 13: Phase 2 - Between Report Date and Issuance
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_background(s13, BG_LIGHT)
    add_accounting_header(s13, "بخش چهارم: روش‌های رسیدگی حسابرسی", "مرحله دوم: حقایق کشف شده پس از تاریخ گزارش و پیش از انتشار", "مسئولیت‌ها، نحوه برخورد با اصلاح مدیریت و اقدامات قانونی حسابرس")
    
    c_p2_1 = create_accounting_card(s13, Inches(6.8), Inches(1.8), w6, h6, ACCENT_EMERALD)
    tf = c_p2_1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"حالت اول: اصلاح صورت‌های مالی توسط مدیریت", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- حسابرس نسبت به اجرای روش{NJ}ها پس از تاریخ گزارش مسئولیتی ندارد؛ مسئولیت آگاه کردن حسابرس با مدیر واحد مورد رسیدگی است.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- در صورت اطلاع از حقایق بااهمیت، حسابرس ضرورت تعدیل را ارزیابی و موضوع را با مدیریت در میان می{NJ}گذارد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"اقدامات حسابرس در صورت اصلاح توسط مدیریت:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_EMERALD, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"۱. اجرای روش{NJ}های حسابرسی مورد نیاز در شرایط موجود\n۲. صدور گزارش جدید درباره صورت{NJ}های مالی اصلاح شده\n۳. تاریخ گزارش جدید نباید پیش از تاریخ تأیید صورت{NJ}های اصلاحی باشد", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    c_p2_2 = create_accounting_card(s13, Inches(0.8), Inches(1.8), w6, h6, ACCENT_ROSE)
    tf = c_p2_2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"حالت دوم: خودداری مدیریت از اصلاح صورت‌ها", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"چنانچه حسابرس به ضرورت اصلاح اعتقاد داشته باشد اما مدیریت از اصلاح صورت{NJ}های مالی خودداری کند:", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"الف) اگر گزارش هنوز تسلیم نشده باشد:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_ROSE, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"حسابرس باید نظر مشروط [Qualified] یا مردود [Adverse] ارائه کند.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(6))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"ب) اگر گزارش قبلاً به واحد ارائه شده باشد:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_ROSE, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- هشدار به مسئولان اصلی اداره امور (هیئت مدیره) مبنی بر خودداری از انتشار صورت{NJ}ها و گزارش\n- انجام اقدامات لازم برای جلوگیری از اتکا به گزارش (اعلام به مجامع سهامداران، بورس، اعتباردهندگان و مراجع قانونی).", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    add_accounting_footer(s13, 13)

    # ==========================================
    # SLIDE 14: Phase 3 - Discovered After Issuance
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    add_background(s14, BG_LIGHT)
    add_accounting_header(s14, "بخش چهارم: روش‌های رسیدگی حسابرسی", "مرحله سوم: حقایق کشف شده پس از انتشار صورت‌های مالی", "رسیدگی به حقایق موجود در تاریخ گزارش که پس از انتشار عمومی افشا شده‌اند")
    
    c_p3_1 = create_accounting_card(s14, Inches(6.8), Inches(1.8), w6, h6, ACCENT_INDIGO)
    tf = c_p3_1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"تجدیدنظر در صورت‌های مالی توسط مدیریت", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- شرط شمول: حقایقی کشف شود که در تاریخ گزارش وجود داشته و اطلاع از آن{NJ}ها موجب تغییر در گزارش حسابرس می{NJ}شده است.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"اقدامات حسابرس در صورت تصمیم به تجدیدنظر:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_INDIGO, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"۱. بررسی اقدامات مدیریت جهت اطمینان از آگاه شدن تمام دریافت{NJ}کنندگان قبلی\n۲. اجرای روش{NJ}های حسابرسی مورد نیاز در آن شرایط\n۳. صدور گزارش جدید شامل یک «بند تأکید بر مطلب خاص»\n۴. اشاره به یادداشت همراه و تشریح علت تجدیدنظر در صورت{NJ}ها", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    c_p3_2 = create_accounting_card(s14, Inches(0.8), Inches(1.8), w6, h6, ACCENT_ROSE)
    tf = c_p3_2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"امتناع مدیریت و استثنای گزارشگری", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"واکنش حسابرس در صورت عدم تجدیدنظر مدیریت:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_ROSE, space_after=Pt(4))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- هشدار رسمی به مسئولان اصلی اداره امور (هیئت مدیره)\n- انجام اقدامات مقتضی جهت جلوگیری از هرگونه اتکا به گزارش در آینده با توجه به اختیارات قانونی و توصیه مشاور حقوقی.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(8))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"استثنای کاربردی در استانداردهای حسابرسی:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=TEXT_DARK, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"در صورت مشخص و نزدیک بودن زمان انتشار صورت{NJ}های مالی دوره بعد، تجدیدنظر در صورت{NJ}های قبلی به شرطی ضرورت ندارد که افشای مناسب موضوع، در صورت{NJ}های دوره بعد انجام گیرد.", font_name=FONT_NAZANIN, font_size=Pt(12), bold=True, color=ACCENT_GOLD)

    add_accounting_footer(s14, 14)

    # ==========================================
    # SLIDE 15: Dual Dating of Audit Report
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    add_background(s15, BG_LIGHT)
    add_accounting_header(s15, "بخش پنجم: تاریخ‌گذاری دوگانه", "تاریخ‌گذاری دوگانه گزارش حسابرسی [Dual Dating]", "رویکردی هوشمندانه برای تحدید مسئولیت قانونی و حرفه‌ای حسابرس")
    
    c_dd_1 = create_accounting_card(s15, Inches(6.8), Inches(1.8), w6, h6, ACCENT_TEAL)
    tf = c_dd_1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"مفهوم و نمونه کاربردی تاریخ‌گذاری دوگانه", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- قاعده عمومی: عموماً تاریخ گزارش حسابرس همان تاریخ پایان اجرای عملیات حسابرسی است و مسئولیت تا آن تاریخ ادامه دارد.", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- سناریو: اگر دعوای حقوقی افشا شده در یادداشت{NJ}ها، در ۱۱ تیر با پرداخت وجه خاتمه یابد و در یادداشت{NJ}ها تعدیل شود:", font_name=FONT_NAZANIN, font_size=Pt(12.5), color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"فرمت استاندارد تاریخ‌گذاری دوگانه:", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_TEAL, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"«۳ تیر ماه، به استثنای یادداشت شماره X که تاریخ آن ۱۲ تیر ماه است.»", font_name=FONT_TITR, font_size=Pt(13.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"این تاریخ{NJ}گذاری مسئولیت حسابرس را تنها برای افشای همان یک قلم خاص تا تاریخ دوم گسترش می{NJ}دهد.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_MUTED)

    c_dd_2 = create_accounting_card(s15, Inches(0.8), Inches(1.8), w6, h6, ACCENT_GOLD)
    tf = c_dd_2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p = tf.paragraphs[0]
    apply_fa_text(p, f"مقایسه راهبردی دو گزینه پیش‌روی حسابرس", font_name=FONT_TITR, font_size=Pt(14.5), bold=True, color=TEXT_DARK, space_after=Pt(6))
    p = tf.add_paragraph()
    apply_fa_text(p, f"گزینه اول: تاریخ‌گذاری دوگانه [Dual Dating]", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_TEAL, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- مزیت کلیدی: مسئولیت حسابرس صرفاً برای افشای همان یک قلم خاص تا دومین تاریخ گسترش می{NJ}یابد.\n- عدم نیاز به تکرار روش{NJ}ها برای سایر اقلام و صرفه{NJ}جویی چشمگیر در زمان و ریسک حرفه{NJ}ای.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK, space_after=Pt(8))
    
    p = tf.add_paragraph()
    apply_fa_text(p, f"گزینه دوم: تمدید تاریخ گزارش برای کل صورت‌ها", font_name=FONT_TITR, font_size=Pt(12.5), bold=True, color=ACCENT_ROSE, space_after=Pt(3))
    p = tf.add_paragraph()
    apply_fa_text(p, f"- استفاده از تاریخ دوم به عنوان تاریخ گزارش (۱۲ تیر):\n- مسئولیت حسابرسان را در قبال تمامی مندرجات صورت{NJ}های مالی گسترش می{NJ}دهد و او باید تمام روش{NJ}های رسیدگی را تجدید کند.", font_name=FONT_NAZANIN, font_size=Pt(12), color=TEXT_DARK)

    add_accounting_footer(s15, 15)

    # ==========================================
    # SLIDE 16: Summary Table / Accounting Matrix
    # ==========================================
    s16 = prs.slides.add_slide(blank_layout)
    add_background(s16, BG_LIGHT)
    add_accounting_header(s16, "بخش ششم: جمع‌بندی و نتیجه‌گیری", "ماتریس مقایسه‌ای و خلاصه تحلیلی رویدادها و وظایف حسابرس", "مرور جامع تفاوت‌های بنیادین، نوع افشا و سطح مسئولیت‌های حرفه‌ای")
    
    table_shape = s16.shapes.add_table(5, 4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.2) # Feature
    table.columns[1].width = Inches(3.2) # Adjusting
    table.columns[2].width = Inches(3.2) # Non-Adjusting
    table.columns[3].width = Inches(3.133) # Auditor Duty
    
    headers = ["شاخص و ویژگی مقایسه", "رویدادهای تعدیلی [گروه اول]", "رویدادهای غیرتعدیلی [گروه دوم]", "وظایف و الزامات حسابرس"]
    for i, h_text in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DEEP
        p = cell.text_frame.paragraphs[0]
        apply_fa_text(p, h_text, font_name=FONT_TITR, font_size=Pt(11.5), bold=True, color=TEXT_LIGHT, align=PP_ALIGN.CENTER, space_after=Pt(0))
        
    rows_data = [
        (f"منشأ و زمان شرایط", f"شواهد درباره شرایط موجود در تاریخ ترازنامه", f"شرایطی که تماماً پس از تاریخ ترازنامه حادث شده", f"تشخیص دقیق ریشه زمانی رویداد با شواهد کافی"),
        (f"نحوه برخورد حسابداری", f"اصلاح و تعدیل ارقام در متن صورت{NJ}های مالی اساسی", f"عدم اصلاح ارقام؛ صرفاً افشا در یادداشت{NJ}های پیوست", f"اطمینان از کفایت افشا یا صحت ثبت سند تعدیلی"),
        (f"تأثیر بر استفاده{NJ}کنندگان", f"برطرف{NJ}کننده ابهام برآوردهای پایان دوره مالی", f"فراهم{NJ}کننده تصویر منصفانه از چشم{NJ}انداز آینده", f"جلوگیری از صدور نظر مقبول در صورت عدم اصلاح"),
        (f"ابزار گزارشگری حسابرس", f"صدور گزارش با ارقام اصلاح{NJ}شده و تاریخ جدید", f"بند تأکید بر مطلب خاص یا تاریخ{NJ}گذاری دوگانه", f"تحدید مسئولیت حرفه{NJ}ای با تاریخ{NJ}گذاری دوگانه")
    ]
    
    for row_idx, r_data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(r_data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 1 else RGBColor(241, 245, 249)
            p = cell.text_frame.paragraphs[0]
            is_bold = (col_idx == 0)
            text_color = ACCENT_INDIGO if col_idx == 0 else TEXT_DARK
            font_use = FONT_TITR if col_idx == 0 else FONT_NAZANIN
            apply_fa_text(p, text, font_name=font_use, font_size=Pt(11.5), bold=is_bold, color=text_color, align=PP_ALIGN.RIGHT, space_after=Pt(0))

    add_accounting_footer(s16, 16)

    # ==========================================
    # SLIDE 17: Closing Slide (Q&A and Gratitude)
    # ==========================================
    s17 = prs.slides.add_slide(blank_layout)
    add_background(s17, NAVY_DEEP)
    
    frame17 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(0.9), Inches(10.333), Inches(5.7))
    frame17.fill.solid()
    frame17.fill.fore_color.rgb = SLATE_DARK
    frame17.line.color.rgb = RGBColor(51, 65, 85)
    frame17.line.width = Pt(1.5)
    
    tf17 = frame17.text_frame
    tf17.word_wrap = True
    
    p = tf17.paragraphs[0]
    apply_fa_text(p, f"با سپاس و قدردانی فراوان", font_name=FONT_TITR, font_size=Pt(32), bold=True, color=TEXT_LIGHT, align=PP_ALIGN.CENTER, space_after=Pt(8))
    
    p = tf17.add_paragraph()
    apply_fa_text(p, f"از توجه و راهنمایی‌های ارزشمند استاد گرانقدر:", font_name=FONT_NAZANIN, font_size=Pt(18), bold=False, color=TEXT_LIGHT_MUTED, align=PP_ALIGN.CENTER, space_after=Pt(4))
    
    p = tf17.add_paragraph()
    apply_fa_text(p, f"جناب استاد بهرامیان", font_name=FONT_TITR, font_size=Pt(28), bold=True, color=RGBColor(253, 224, 71), align=PP_ALIGN.CENTER, space_after=Pt(6))
    
    p = tf17.add_paragraph()
    apply_fa_text(p, f"ارائه دهنده: فاطمه عیوضی", font_name=FONT_TITR, font_size=Pt(20), bold=True, color=RGBColor(52, 211, 153), align=PP_ALIGN.CENTER, space_after=Pt(10))
    
    p = tf17.add_paragraph()
    apply_fa_text(p, f"________________________", font_name=FONT_NAZANIN, font_size=Pt(16), bold=False, color=RGBColor(71, 85, 105), align=PP_ALIGN.CENTER, space_after=Pt(10))
    
    p = tf17.add_paragraph()
    apply_fa_text(p, f"بخش پرسش و پاسخ [Q & A]", font_name=FONT_TITR, font_size=Pt(20), bold=True, color=RGBColor(56, 189, 248), align=PP_ALIGN.CENTER, space_after=Pt(4))
    
    p = tf17.add_paragraph()
    apply_fa_text(p, f"مشتاق شنیدن نظرات، دیدگاه‌ها و سوالات شما عزیزان هستیم", font_name=FONT_NAZANIN, font_size=Pt(14), bold=False, color=TEXT_LIGHT_MUTED, align=PP_ALIGN.CENTER)

    # Save presentations
    prs.save(output_path)
    eng_path = os.path.join(os.path.dirname(output_path), "Events_After_Balance_Sheet_Date.pptx")
    prs.save(eng_path)
    print("Accounting presentation created successfully!")

if __name__ == '__main__':
    out_dir = r"c:\Users\Sadeghi\OneDrive\Desktop\New folder (7)"
    out_file = os.path.join(out_dir, "رویدادهای_بعد_از_تاریخ_ترازنامه.pptx")
    create_accounting_presentation(out_file)
